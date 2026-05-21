from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .import_excel import REQUIRED_COLUMNS
from .models import GROUP_NAMES, SUBJECT_NAMES


def style_header(sheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font


def autosize_columns(sheet) -> None:
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 70)


def build_recommendations_workbook(recommendations: list[dict]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recommendations"

    headers = [
        "Nama Siswa",
        "NIS",
        "Kelas Asal",
        *[f"Skor {group}" for group in GROUP_NAMES],
        "Recommended Group",
        "Recommended Rombel",
        "Alternative Rombels",
        "Final Group",
        "Final Rombel",
        "Status",
        "Review Notes",
        "Override",
    ]
    sheet.append(headers)

    style_header(sheet)

    for item in recommendations:
        scores = item.get("scoresByGroup", {})
        sheet.append(
            [
                item.get("studentName", ""),
                item.get("nis", ""),
                item.get("originClass", ""),
                *[scores.get(group, 0) for group in GROUP_NAMES],
                item.get("recommendedGroup", ""),
                item.get("recommendedRombel") or "",
                ", ".join(item.get("alternativeRombels", [])),
                item.get("finalGroup") or "",
                item.get("finalRombel") or "",
                item.get("status", ""),
                item.get("reviewNotes", ""),
                "Yes" if item.get("isOverridden") else "No",
            ]
        )

    autosize_columns(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_students_workbook(students: list[dict]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    headers = [
        "Nama",
        "NIS",
        "Kelas Asal",
        "Cita-cita",
        "Rencana/Jurusan Tujuan",
        "Mapel Pilihan 1",
        "Mapel Pilihan 2",
        "Mapel Cadangan",
        "Mapel Terkuat",
        *[f"Skor {subject}" for subject in SUBJECT_NAMES],
        "Alasan/Catatan",
        "Created At",
        "Updated At",
    ]
    sheet.append(headers)
    style_header(sheet)

    for student in students:
        scores = student.get("scores", {})
        sheet.append(
            [
                student.get("name", ""),
                student.get("nis", ""),
                student.get("originClass", ""),
                student.get("careerGoal", ""),
                student.get("targetMajor", ""),
                student.get("prioritySubject1", ""),
                student.get("prioritySubject2", ""),
                student.get("backupSubject", ""),
                student.get("strongestSubject", ""),
                *[scores.get(subject, 0) for subject in SUBJECT_NAMES],
                student.get("reason", ""),
                student.get("createdAt", ""),
                student.get("updatedAt", ""),
            ]
        )
    autosize_columns(sheet)

    import_sheet = workbook.create_sheet("2026")
    import_sheet.append(REQUIRED_COLUMNS)
    style_header(import_sheet)
    for index, student in enumerate(students, start=1):
        import_sheet.append(
            [
                student.get("nis") or index,
                student.get("name", ""),
                student.get("originClass", ""),
                student.get("careerGoal", ""),
                student.get("prioritySubject1", ""),
                student.get("prioritySubject2", ""),
                student.get("backupSubject", ""),
                student.get("strongestSubject", ""),
                "",
                student.get("targetMajor", ""),
            ]
        )
    autosize_columns(import_sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

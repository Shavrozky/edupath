from datetime import datetime
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from .models import SUBJECT_NAMES
from .scoring import now_iso


SUBJECT_ALIASES = {
    "BIOLOGI": "Biologi",
    "KIMIA": "Kimia",
    "FISIKA": "Fisika",
    "MATEMATIKA TL": "Matematika Tindak Lanjut",
    "MATEMATIKA TINDAK LANJUT": "Matematika Tindak Lanjut",
    "EKONOMI": "Ekonomi",
    "SOSIOLOGI": "Sosiologi",
    "GEOGRAFI": "Geografi",
    "BAHASA INGGRIS": "Bahasa Inggris Tindak Lanjut",
    "BAHASA INGGRIS TINDAK LANJUT": "Bahasa Inggris Tindak Lanjut",
    "BAHASA ARAB": "Bahasa Arab",
}

REQUIRED_COLUMNS = [
    "No..",
    "Nama Lengkap ",
    "Kelas",
    "Apa profesi yang kamu cita-citakan? Mengapa kamu memilih profesi tersebut?\nKamu bisa menuliskan lebih dari 1",
    "Mata Pelajaran Pilihan 1",
    "Mata Pelajaran Pilihan 2",
    "Mata Pelajaran Pilihan 3",
    "Mata Pelajaran Pilihan 4",
    "Mata Pelajaran Pilihan 5",
    "Apa Kira-kira yang akan kamu rencanakan setelah lulus SMA. \nJika kalian memilih bekerja maka pekerjaan apa yang akan kamu pilih. \njika kalian memilih untuk melanjutkan perkuliahan maka dalam bidang apa kamu ingin berkuliah?\nberikan jawaban beserta alasan kalian.",
]


def normalize_subject(value: Any) -> str | None:
    if value is None:
        return None
    key = str(value).strip().upper()
    return SUBJECT_ALIASES.get(key)


def default_scores() -> dict[str, float]:
    return {subject: 0.0 for subject in SUBJECT_NAMES}


def build_reason(career_goal: str, plan: str, extra_subjects: list[str]) -> str:
    parts = []
    if career_goal:
        parts.append(f"Cita-cita: {career_goal}")
    if plan:
        parts.append(f"Rencana setelah SMA: {plan}")
    if extra_subjects:
        parts.append(f"Pilihan mapel tambahan: {', '.join(extra_subjects)}")
    return " | ".join(parts)


def read_students_from_excel(content: bytes, sheet_name: str = "2026") -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' tidak ditemukan. Sheet tersedia: {', '.join(workbook.sheetnames)}")

    sheet = workbook[sheet_name]
    headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Kolom wajib tidak ditemukan: {', '.join(missing)}")

    indexes = {header: headers.index(header) for header in REQUIRED_COLUMNS}
    students: list[dict[str, Any]] = []
    warnings: list[str] = []
    timestamp = now_iso()

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = str(row[indexes["Nama Lengkap "]] or "").strip()
        if not name:
            continue
        if name.lower() == "nama lengkap":
            continue

        number = row[indexes["No.."]]
        nis = str(number).strip() if number is not None else str(row_number - 1)
        origin_class = str(row[indexes["Kelas"]] or "").strip()
        career_goal = str(row[indexes["Apa profesi yang kamu cita-citakan? Mengapa kamu memilih profesi tersebut?\nKamu bisa menuliskan lebih dari 1"]] or "").strip()
        plan = str(row[indexes["Apa Kira-kira yang akan kamu rencanakan setelah lulus SMA. \nJika kalian memilih bekerja maka pekerjaan apa yang akan kamu pilih. \njika kalian memilih untuk melanjutkan perkuliahan maka dalam bidang apa kamu ingin berkuliah?\nberikan jawaban beserta alasan kalian."]] or "").strip()

        choices = []
        for choice_number in range(1, 6):
            raw_subject = row[indexes[f"Mata Pelajaran Pilihan {choice_number}"]]
            if raw_subject is None or str(raw_subject).strip() == "":
                continue
            subject = normalize_subject(raw_subject)
            if subject is None:
                warnings.append(f"Baris {row_number}: mapel pilihan {choice_number} tidak dikenali: {raw_subject!r}")
            else:
                choices.append(subject)

        if not choices:
            warnings.append(f"Baris {row_number}: dilewati karena tidak ada mapel valid untuk {name}")
            continue

        priority_subject_1 = choices[0]
        priority_subject_2 = choices[1] if len(choices) > 1 else priority_subject_1
        backup_subject = choices[2] if len(choices) > 2 else priority_subject_2
        strongest_subject = priority_subject_1
        extra_subjects = choices[3:]

        students.append(
            {
                "id": f"student-{uuid4()}",
                "name": name,
                "nis": nis,
                "originClass": origin_class,
                "careerGoal": career_goal,
                "targetMajor": plan,
                "prioritySubject1": priority_subject_1,
                "prioritySubject2": priority_subject_2,
                "backupSubject": backup_subject,
                "strongestSubject": strongest_subject,
                "scores": default_scores(),
                "reason": build_reason(career_goal, plan, extra_subjects),
                "createdAt": timestamp,
                "updatedAt": timestamp,
            }
        )

    return students, warnings


def merge_students(existing: list[dict[str, Any]], imported: list[dict[str, Any]], replace: bool) -> list[dict[str, Any]]:
    if replace:
        return imported

    by_nis = {student["nis"]: student for student in existing}
    for student in imported:
        by_nis[student["nis"]] = student
    return list(by_nis.values())
